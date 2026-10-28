#!/usr/bin/env python3
"""
FIR feature extraction and dataset builder.

Supports:
1. Single-run mode using individual command-line arguments.
2. Multi-run mode using a CSV manifest.

Reads existing OpenROAD results; does not rerun physical design.

Requires:
    extract_fir_features_readable.py
    Python 3.10+
    openpyxl (optional, for Excel output)
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from pathlib import Path
from typing import Any

import extract_fir_features_readable as base


# ---------------------------------------------------------------------
# General helpers
# ---------------------------------------------------------------------

def read_manifest(path: Path) -> list[dict[str, str]]:
    """Read and validate the run manifest."""
    if not path.is_file():
        raise FileNotFoundError(f"Manifest CSV not found: {path}")

    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)

        if not reader.fieldnames:
            raise ValueError("Manifest CSV has no header.")

        required = {"run_dir", "netlist_json"}
        missing = required - set(reader.fieldnames)

        if missing:
            raise ValueError(
                "Manifest is missing required columns: "
                + ", ".join(sorted(missing))
            )

        rows = []

        for line_number, row in enumerate(reader, start=2):
            if not row.get("run_dir", "").strip():
                raise ValueError(
                    f"Manifest line {line_number}: run_dir is empty."
                )

            if not row.get("netlist_json", "").strip():
                raise ValueError(
                    f"Manifest line {line_number}: netlist_json is empty."
                )

            rows.append({
                str(key).strip(): (
                    value.strip() if isinstance(value, str) else ""
                )
                for key, value in row.items()
                if key is not None
            })

    if not rows:
        raise ValueError("Manifest contains no runs.")

    return rows


def optional_int(row: dict[str, str], key: str) -> int | None:
    value = row.get(key, "").strip()

    if not value:
        return None

    result = int(value)

    if result <= 0:
        raise ValueError(f"{key} must be a positive integer.")

    return result


def optional_float(row: dict[str, str], key: str) -> float | None:
    value = row.get(key, "").strip()

    if not value:
        return None

    result = float(value)

    if not math.isfinite(result):
        raise ValueError(f"{key} must be a finite number.")

    if key == "utilization" and not 0 < result <= 100:
        raise ValueError("utilization must be greater than 0 and <= 100.")

    if key == "place_density" and not 0 < result <= 1:
        raise ValueError("place_density must be greater than 0 and <= 1.")

    if key == "clock_period" and result <= 0:
        raise ValueError("clock_period must be positive.")

    return result


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    with path.open("w", encoding="utf-8") as handle:
        json.dump(obj, handle, indent=2, allow_nan=False)
        handle.write("\n")


def write_dict_csv(path: Path, row: dict[str, Any]) -> None:
    """Write one record, converting nested values to JSON."""
    path.parent.mkdir(parents=True, exist_ok=True)

    flattened = {}

    for key, value in row.items():
        if isinstance(value, (dict, list, tuple)):
            flattened[key] = json.dumps(value, ensure_ascii=False)
        else:
            flattened[key] = value

    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(flattened))
        writer.writeheader()
        writer.writerow(flattened)


def write_rows_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    """Write multiple records with a stable union of column names."""
    path.parent.mkdir(parents=True, exist_ok=True)

    if not rows:
        path.write_text("", encoding="utf-8")
        return

    fields = []
    seen = set()

    for row in rows:
        for key in row:
            if key not in seen:
                fields.append(key)
                seen.add(key)

    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=fields,
            extrasaction="ignore",
        )

        writer.writeheader()

        for row in rows:
            clean = {}

            for key in fields:
                value = row.get(key)

                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(value, ensure_ascii=False)

                clean[key] = value

            writer.writerow(clean)


def safe_name(value: str) -> str:
    """Create a filesystem-safe directory name."""
    result = "".join(
        char if char.isalnum() or char in "-_." else "_"
        for char in value
    )

    return result.strip("._") or "run"


def markdown_escape(value: Any) -> str:
    text = base.display_value(value)
    return text.replace("|", r"\|").replace("\n", " ")


def markdown_table(rows: list[tuple[str, Any]]) -> str:
    lines = ["| Metric | Value |", "|---|---|"]

    for label, value in rows:
        lines.append(
            f"| {markdown_escape(label)} | {markdown_escape(value)} |"
        )

    return "\n".join(lines)


# ---------------------------------------------------------------------
# Fanout metadata coverage
# ---------------------------------------------------------------------

def fanout_metadata_coverage(
    netlist_path: Path,
    top_module: str,
) -> dict[str, Any]:
    """Measure availability of port-direction metadata."""
    design = base.load_netlist(netlist_path)
    modules = design["modules"]

    if top_module not in modules:
        top_module = base.select_top(modules, top_module)

    cells = modules[top_module].get("cells", {})

    total_cells = len(cells)
    cells_with_directions = 0
    cells_with_complete_directions = 0

    for cell in cells.values():
        connections = cell.get("connections", {})
        directions = cell.get("port_directions", {})

        if isinstance(directions, dict) and directions:
            cells_with_directions += 1

        if (
            isinstance(connections, dict)
            and isinstance(directions, dict)
            and connections
            and all(
                port in directions
                and str(directions[port]).lower()
                in {"input", "output", "inout"}
                for port in connections
            )
        ):
            cells_with_complete_directions += 1

    coverage = (
        cells_with_complete_directions / total_cells
        if total_cells
        else None
    )

    return {
        "fanout_total_cells": total_cells,
        "fanout_cells_with_any_directions": cells_with_directions,
        "fanout_cells_with_complete_directions":
            cells_with_complete_directions,
        "fanout_direction_coverage_fraction": coverage,
        "fanout_direction_coverage_percent": (
            coverage * 100 if coverage is not None else None
        ),
        "fanout_metadata_complete_for_all_cells": (
            total_cells > 0
            and cells_with_complete_directions == total_cells
        ),
    }


# ---------------------------------------------------------------------
# Per-run extraction
# ---------------------------------------------------------------------

def extract_one(
    row: dict[str, str],
    output_root: Path,
    manifest_dir: Path,
    output_dir: Path | None = None,
) -> dict[str, Any]:
    """Extract synthesis and physical-design features for one run."""

    def resolve_path(value: str) -> Path:
        path = Path(value).expanduser()

        if not path.is_absolute():
            path = manifest_dir / path

        return path.resolve()

    run_dir = resolve_path(row["run_dir"])
    netlist_path = resolve_path(row["netlist_json"])

    run_id = row.get("run_id") or run_dir.name
    design_id = row.get("design_id") or run_id
    requested_top = row.get("top") or "fir_filter"

# Direct mode can write to the existing FIR run directory.
# Manifest mode retains a separate output folder for each run.
    run_output = (
        output_dir.resolve()
        if output_dir is not None
        else output_root / safe_name(run_id)
    )
    run_output.mkdir(parents=True, exist_ok=True)

    # Extract synthesis/netlist features.
    pre = base.extract_netlist_features(
        netlist_path,
        requested_top,
        design_id,
    )

    actual_top = pre.get("top_module", requested_top)

    pre["run_id"] = run_id
    pre["design_id"] = design_id

    pre["utilization_constraint_percent"] = optional_float(
        row, "utilization"
    )
    pre["clock_period_ns"] = optional_float(row, "clock_period")

    pre["taps"] = optional_int(row, "taps")
    pre["data_width"] = optional_int(row, "data_width")
    pre["coeff_width"] = optional_int(row, "coeff_width")
    pre["acc_width"] = optional_int(row, "acc_width")
    pre["place_density"] = optional_float(row, "place_density")

    # Read existing physical-design artifacts; do not rerun OpenROAD.
    post = base.physical_features(run_dir)

    coverage = fanout_metadata_coverage(netlist_path, actual_top)
    pre.update(coverage)

    if not pre["fanout_metadata_complete_for_all_cells"]:
        pre["fanout_quality_warning"] = (
            "Some cells lack complete port-direction metadata. "
            "Fanout statistics may be incomplete."
        )
    else:
        pre["fanout_quality_warning"] = (
            "All cells have complete port-direction metadata."
        )

    # Replace non-finite floats before JSON serialization.
    for mapping in (pre, post):
        for key, value in list(mapping.items()):
            if isinstance(value, float) and not math.isfinite(value):
                mapping[key] = None

    combined = {**pre, **post}

    # Determine observed flow status.
    def_ok = bool(
        post.get("final_def_available")
        and post.get("final_def_nonempty")
    )
    odb_ok = bool(post.get("final_odb_available"))
    route_success = bool(
        post.get("detailed_route_success_marker_found")
    )

    if route_success and def_ok and odb_ok:
        flow_status = "detailed_routing_completed"
        last_stage = "detailed_routing"
    elif def_ok and odb_ok:
        flow_status = "artifacts_present"
        last_stage = "unknown"
    elif def_ok or odb_ok:
        flow_status = "partial_artifacts"
        last_stage = "unknown"
    elif post.get("log_files_found"):
        flow_status = "flow_incomplete_or_unverified"
        last_stage = "unknown"
    else:
        flow_status = "no_artifacts_or_logs"
        last_stage = "unknown"

    config = {
        "run_id": run_id,
        "design_id": design_id,
        "top_module": actual_top,
        "netlist_json": str(netlist_path),
        "run_dir": str(run_dir),
        "taps": pre["taps"],
        "data_width": pre["data_width"],
        "coeff_width": pre["coeff_width"],
        "acc_width": pre["acc_width"],
        "utilization": pre["utilization_constraint_percent"],
        "place_density": pre["place_density"],
        "clk_period_ns": pre["clock_period_ns"],
        "flow_status": flow_status,
        "last_completed_stage": last_stage,
    }

    # Individual-run outputs.
    write_json(run_output / "preplacement_features.json", pre)
    write_dict_csv(run_output / "preplacement_features.csv", pre)

    write_json(run_output / "physical_targets.json", post)
    write_dict_csv(run_output / "physical_targets.csv", post)

    write_json(run_output / "features.json", combined)
    write_dict_csv(run_output / "features.csv", combined)

    write_json(
        run_output / "extraction_status.json",
        {
            "run_id": run_id,
            "design_id": design_id,
            "top_module": actual_top,
            "run_dir": str(run_dir),
            "netlist_json": str(netlist_path),
            "output_dir": str(run_output),
            "flow_status": flow_status,
            "fanout_direction_coverage_percent": coverage[
                "fanout_direction_coverage_percent"
            ],
            "fanout_metadata_complete_for_all_cells": coverage[
                "fanout_metadata_complete_for_all_cells"
            ],
            "target_measurement_status": post.get(
                "target_measurement_status", {}
            ),
        },
    )

    # Human-readable Markdown report.
    config_rows = [
        ("Run ID", run_id),
        ("Design ID", design_id),
        ("Top module selected", actual_top),
        ("Netlist JSON", str(netlist_path)),
        ("Run directory", str(run_dir)),
        ("FIR taps", pre["taps"]),
        ("Data width (bits)", pre["data_width"]),
        ("Coefficient width (bits)", pre["coeff_width"]),
        ("Accumulator width (bits)", pre["acc_width"]),
        ("Utilization constraint (%)",
         pre["utilization_constraint_percent"]),
        ("Placement density", pre["place_density"]),
        ("Clock period (ns)", pre["clock_period_ns"]),
        ("Detected flow status", flow_status),
    ]

    synthesis_keys = [
        "cell_count",
        "netname_count",
        "port_count",
        "port_bit_count",
        "sequential_cell_count",
        "combinational_cell_count",
        "buffer_inverter_count",
        "other_cell_count",
        "unique_cell_type_count",
        "avg_bit_fanout",
        "max_bit_fanout",
        "std_bit_fanout",
        "fanout_direction_coverage_percent",
        "fanout_metadata_complete_for_all_cells",
        "fanout_quality_warning",
        "total_net_bits",
        "cell_classification_method",
    ]

    target_keys = [
        "final_def_available",
        "final_odb_available",
        "final_def_nonempty",
        "artifacts_present",
        "detailed_route_success_marker_found",
        "routed_net_count",
        "route_segment_count",
        "wirelength_um",
        "via_count",
        "final_detailed_route_violations",
        "effective_place_density",
        "design_area_um2",
        "wns_setup_ns",
        "tns_setup_ns",
        "wns_hold_ns",
        "congestion_overflow",
    ]

    source_keys = [
        "setup_report_source",
        "tns_report_source",
        "hold_report_source",
        "congestion_report_source",
        "routing_log_source",
        "target_feature_stage",
        "note",
    ]

    lines = [
        "# FIR Physical-Design Feature Report",
        "",
        f"**Run:** `{run_id}`",
        "",
        "## 1. Configuration",
        "",
        markdown_table(config_rows),
        "",
        "## 2. Synthesis and netlist features",
        "",
        markdown_table([
            (
                base.DISPLAY_NAMES.get(
                    key, key.replace("_", " ").title()
                ),
                pre.get(key),
            )
            for key in synthesis_keys
        ]),
        "",
        "## 3. Physical-design results",
        "",
        markdown_table([
            (
                base.DISPLAY_NAMES.get(
                    key, key.replace("_", " ").title()
                ),
                post.get(key),
            )
            for key in target_keys
        ]),
        "",
        "## 4. Measurement sources and notes",
        "",
        markdown_table([
            (
                base.DISPLAY_NAMES.get(
                    key, key.replace("_", " ").title()
                ),
                post.get(key),
            )
            for key in source_keys
        ]),
        "",
        "## 5. Interpretation notes",
        "",
        "- N/A means unavailable, not zero.",
        "- DEF route counts are approximate.",
        "- Cell classification is heuristic and should be validated.",
        "- Fanout metrics require complete port-direction metadata.",
        "- Avoid target leakage when preparing ML predictors.",
        "- Detailed-route violations are not independent signoff DRC.",
        "- Validate timing extraction against the actual report format.",
        "",
        "## 6. Files written for this run",
        "",
        "- `preplacement_features.csv` and `.json`",
        "- `physical_targets.csv` and `.json`",
        "- `features.csv` and `.json`",
        "- `extraction_status.json`",
        "- `run_report.md`",
        "",
    ]

    (run_output / "run_report.md").write_text(
        "\n".join(lines),
        encoding="utf-8",
    )

    return {
        "config": config,
        "pre": pre,
        "post": post,
        "combined": combined,
        "output_dir": str(run_output),
    }


# ---------------------------------------------------------------------
# Excel workbook
# ---------------------------------------------------------------------

def write_workbook(
    output_root: Path,
    results: list[dict[str, Any]],
) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Font,
            PatternFill,
            Alignment,
            Border,
            Side,
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(
            "WARNING: Excel workbook skipped. Install openpyxl with "
            "'python -m pip install openpyxl'.",
            file=sys.stderr,
        )
        return False

    wb = Workbook()

    ws_config = wb.active
    ws_config.title = "Configuration"
    ws_synth = wb.create_sheet("Synthesis Features")
    ws_targets = wb.create_sheet("Physical Targets")
    ws_combined = wb.create_sheet("Combined Data")
    ws_dictionary = wb.create_sheet("Metric Dictionary")

    def write_table(
        ws: Any,
        rows: list[dict[str, Any]],
    ) -> None:
        if not rows:
            ws.append(["No data"])
            return

        columns = []
        seen = set()

        for row in rows:
            for key in row:
                if key not in seen:
                    columns.append(key)
                    seen.add(key)

        ws.append(columns)

        for row in rows:
            values = []

            for key in columns:
                value = row.get(key)

                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(value, ensure_ascii=False)

                values.append(value)

            ws.append(values)

    write_table(ws_config, [r["config"] for r in results])
    write_table(ws_synth, [r["pre"] for r in results])
    write_table(ws_targets, [r["post"] for r in results])
    write_table(ws_combined, [r["combined"] for r in results])

    dictionary_rows = [
        {
            "field": "run_id",
            "type": "identifier",
            "unit": "none",
            "stage": "configuration",
            "description": "Unique identifier for this run.",
            "ml_note": "Use for tracking, not normally as a predictor.",
        },
        {
            "field": "cell_count",
            "type": "feature",
            "unit": "cells",
            "stage": "post-synthesis / pre-placement",
            "description": "Number of cells in the selected top module.",
            "ml_note": "Potential predictor.",
        },
        {
            "field": "avg_bit_fanout",
            "type": "feature; validate coverage",
            "unit": "loads per bit",
            "stage": "post-synthesis / pre-placement",
            "description": "Mean sink count over the bits counted.",
            "ml_note": "Use cautiously if direction metadata is incomplete.",
        },
        {
            "field": "wns_setup_ns",
            "type": "target",
            "unit": "ns",
            "stage": "timing analysis",
            "description": "Worst setup slack extracted from timing data.",
            "ml_note": "Exclude when predicting setup WNS.",
        },
        {
            "field": "tns_setup_ns",
            "type": "target",
            "unit": "ns",
            "stage": "timing analysis",
            "description": "Total setup negative slack.",
            "ml_note": "Exclude when predicting setup TNS.",
        },
        {
            "field": "wns_hold_ns",
            "type": "target",
            "unit": "ns",
            "stage": "timing analysis",
            "description": "Worst hold slack extracted from timing data.",
            "ml_note": "Exclude when predicting hold WNS.",
        },
        {
            "field": "wirelength_um",
            "type": "target",
            "unit": "µm",
            "stage": "routing",
            "description": "Wirelength extracted from routing results.",
            "ml_note": "Exclude when predicting wirelength.",
        },
        {
            "field": "via_count",
            "type": "target",
            "unit": "vias",
            "stage": "routing",
            "description": "Via count extracted from routing results.",
            "ml_note": "Exclude when predicting via count.",
        },
        {
            "field": "congestion_overflow",
            "type": "target",
            "unit": "flow-defined",
            "stage": "global routing / congestion",
            "description": "Overflow when a recognized metric is available.",
            "ml_note": "Missing means not extracted, not zero.",
        },
        {
            "field": "design_area_um2",
            "type": "metric; verify report definition",
            "unit": "µm²",
            "stage": "physical-design log",
            "description": "Area parsed from recognized log text.",
            "ml_note": "Verify its measurement stage before using it.",
        },
    ]

    write_table(ws_dictionary, dictionary_rows)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )
                cell.border = Border(bottom=thin)

        ws.row_dimensions[1].height = 30

        for col_cells in ws.columns:
            letter = get_column_letter(col_cells[0].column)

            max_length = max(
                (
                    len(str(cell.value))
                    if cell.value is not None
                    else 0
                )
                for cell in col_cells
            )

            ws.column_dimensions[letter].width = min(
                max(max_length + 2, 14),
                42,
            )

    wb.save(output_root / "fir_multi_run_dataset.xlsx")
    return True


# ---------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Extract FIR features from a single existing run or "
            "multiple runs specified in a CSV manifest."
        )
    )

    # Multi-run options
    parser.add_argument(
        "--manifest",
        default="fir_runs_manifest.csv",
        help="CSV manifest containing one row per design/run.",
    )
    parser.add_argument(
        "--outdir",
        default="fir_multi_run_output",
        help="Output directory for reports and datasets.",
    )

    # Single-run options used by the existing GitHub Actions workflow
    parser.add_argument("--netlist-json")
    parser.add_argument("--run-dir")
    parser.add_argument("--design-id")
    parser.add_argument("--run-id")
    parser.add_argument("--top", default="fir_filter")
    parser.add_argument("--taps")
    parser.add_argument("--data-width")
    parser.add_argument("--coeff-width")
    parser.add_argument("--acc-width")
    parser.add_argument("--utilization")
    parser.add_argument("--place-density")
    parser.add_argument("--clock-period")

    args = parser.parse_args()

    output_root = Path(args.outdir).expanduser().resolve()

    try:
        # Select the input mode.
        direct_mode = (
            args.netlist_json is not None
            or args.run_dir is not None
        )

        if direct_mode:
            if not args.netlist_json or not args.run_dir:
                parser.error(
                    "Single-run mode requires both "
                    "--netlist-json and --run-dir."
                )

            run_id = args.run_id or Path(args.run_dir).name

            rows = [{
                "run_id": run_id,
                "design_id": args.design_id or run_id,
                "run_dir": args.run_dir,
                "netlist_json": args.netlist_json,
                "top": args.top,
                "taps": args.taps or "",
                "data_width": args.data_width or "",
                "coeff_width": args.coeff_width or "",
                "acc_width": args.acc_width or "",
                "utilization": args.utilization or "",
                "place_density": args.place_density or "",
                "clock_period": args.clock_period or "",
            }]

            # Paths supplied by the existing workflow are relative to
            # the repository working directory.
            manifest_dir = Path.cwd()

            print("Input mode: single FIR run")
            print(f"Run ID: {run_id}")

        else:
            manifest_path = (
                Path(args.manifest).expanduser().resolve()
            )

            rows = read_manifest(manifest_path)
            manifest_dir = manifest_path.parent

            print("Input mode: multi-run manifest")
            print(f"Manifest: {manifest_path}")

        output_root.mkdir(parents=True, exist_ok=True)

        results = []
        failures = []

       
        for index, row in enumerate(rows, start=1):
            run_id = row.get("run_id") or Path(row["run_dir"]).name

            print(f"[{index}/{len(rows)}] Extracting {run_id} ...")

            try:
                if direct_mode:
                    # Direct mode: write outputs to the existing run directory.
                    run_dir = Path(row["run_dir"]).expanduser()

                    if not run_dir.is_absolute():
                        run_dir = manifest_dir / run_dir

                    run_dir = run_dir.resolve()

                    result = extract_one(
                        row=row,
                        output_root=output_root,
                        manifest_dir=manifest_dir,
                        output_dir=run_dir,
                    )
                else:
                    # Manifest mode: use a separate directory for each run.
                    result = extract_one(
                        row=row,
                        output_root=output_root / "runs",
                        manifest_dir=manifest_dir,
                    )

                # Handle successful extraction in both modes.
                results.append(result)
                print(f"    OK: {result['output_dir']}")

            except Exception as exc:
                failures.append({
                    "run_id": run_id,
                    "run_dir": row.get("run_dir"),
                    "error": str(exc),
                })

                print(
                    f"    ERROR: {run_id}: {exc}",
                    file=sys.stderr,
                )

        if not results:
            write_json(
                output_root / "failures.json",
                {"failures": failures},
            )

            print(
                "No runs were extracted successfully.",
                file=sys.stderr,
            )
            return 1

        # Consolidated CSV files.
        write_rows_csv(
            output_root / "combined_features_targets.csv",
            [result["combined"] for result in results],
        )

        write_rows_csv(
            output_root / "combined_configurations.csv",
            [result["config"] for result in results],
        )

        write_rows_csv(
            output_root / "combined_synthesis_features.csv",
            [result["pre"] for result in results],
        )

        write_rows_csv(
            output_root / "combined_physical_targets.csv",
            [result["post"] for result in results],
        )

        xlsx_created = write_workbook(output_root, results)

        write_json(
            output_root / "multi_run_status.json",
            {
                "input_mode": (
                    "single_run" if direct_mode else "manifest"
                ),
                "output_dir": str(output_root),
                "runs_requested": len(rows),
                "runs_succeeded": len(results),
                "runs_failed": len(failures),
                "excel_workbook_created": xlsx_created,
                "failures": failures,
                "successful_run_output_directories": [
                    result["output_dir"] for result in results
                ],
                "important_notes": [
                    "Existing artifacts are read; OpenROAD is not rerun.",
                    "N/A/null means unavailable, not zero.",
                    "Validate timing extraction against the actual report.",
                    "Avoid target leakage when preparing ML predictors.",
                ],
            },
        )

        # Consolidated Markdown report.
        report_lines = [
            "# FIR Feature Extraction Report",
            "",
            f"- Input mode: {'single run' if direct_mode else 'manifest'}",
            f"- Runs requested: {len(rows)}",
            f"- Runs extracted successfully: {len(results)}",
            f"- Runs failed: {len(failures)}",
            "",
            "## Run summary",
            "",
            "| Run ID | Design ID | Top module | Flow status | Output directory |",
            "|---|---|---|---|---|",
        ]

        for result in results:
            config = result["config"]

            report_lines.append(
                f"| {markdown_escape(config.get('run_id'))} "
                f"| {markdown_escape(config.get('design_id'))} "
                f"| {markdown_escape(config.get('top_module'))} "
                f"| {markdown_escape(config.get('flow_status'))} "
                f"| {markdown_escape(result.get('output_dir'))} |"
            )

        report_lines.extend([
            "",
            "## Consolidated files",
            "",
            "- `combined_features_targets.csv`",
            "- `combined_configurations.csv`",
            "- `combined_synthesis_features.csv`",
            "- `combined_physical_targets.csv`",
            "- `fir_multi_run_dataset.xlsx` (if openpyxl is installed)",
            "- `multi_run_status.json`",
            "",
            "## Interpretation notes",
            "",
            "- Missing metrics are not zero.",
            "- Routed-net and route-segment counts are approximate.",
            "- Avoid target leakage when selecting ML predictors.",
            "- Validate timing extraction against the actual report format.",
            "",
        ])

        (output_root / "multi_run_report.md").write_text(
            "\n".join(report_lines),
            encoding="utf-8",
        )

        print("\n===== FIR extraction summary =====")
        print(f"Runs requested : {len(rows)}")
        print(f"Runs succeeded : {len(results)}")
        print(f"Runs failed    : {len(failures)}")
        print(f"Output folder  : {output_root}")

        print("\nGenerated files:")

        for filename in (
            "combined_features_targets.csv",
            "combined_configurations.csv",
            "combined_synthesis_features.csv",
            "combined_physical_targets.csv",
            "fir_multi_run_dataset.xlsx",
            "multi_run_report.md",
            "multi_run_status.json",
        ):
            path = output_root / filename

            if path.exists():
                print(f"  - {path}")

        if failures:
            print("\nFailed runs:")

            for failure in failures:
                print(
                    f"  - {failure['run_id']}: {failure['error']}"
                )

        return 0 if not failures else 2

    except (OSError, ValueError, csv.Error) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
